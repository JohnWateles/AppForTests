from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
import os
import win32com.client as win32
import dearpygui.dearpygui as dpg
import pythoncom


def create_excel_table(test_name: str = "example", result: dict = None):
    if result is None:
        result = dict()

    wb = Workbook()
    global_table = wb.active

    suffix = str(len(tuple(name for name in os.listdir(f"results") if test_name in name)) + 1)

    data = [("№", " ", "Да", "Нет")]
    with open(f"tests/{test_name}/questions.txt", encoding="UTF-8") as file:
        questions = [question.rstrip("\n") for question in file.readlines()]
        for i in range(1, len(questions)):
            if result[str(i)] == 0:
                row = (i, questions[i], " ", "+")
            elif result[str(i)] == 1:
                row = (i, questions[i], "+", " ")
            else:
                row = (i, questions[i], " ", " ")
            data.append(row)

    for row in data:
        global_table.append(row)

    table_range = f"A1:D{len(data)}"
    table = Table(displayName="Table1", ref=table_range)
    style = TableStyleInfo(
        name="TableStyleLight8",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=True
    )
    table.tableStyleInfo = style
    global_table.add_table(table)

    for row in global_table.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    global_table.column_dimensions['A'].width = 3
    global_table.column_dimensions['B'].width = 73
    global_table.column_dimensions['C'].width = 5
    global_table.column_dimensions['D'].width = 5

    wb.save(f"results/{test_name}_{suffix}.xlsx")


def print_excel_file():

    file_name = dpg.get_value("file_selection")
    if file_name != "":
        file_path = os.getcwd()
        work_path = [str(elem) for elem in file_path.split("\\")]
        work_path.append("results")
        work_path.append(file_name)
        work_path = "\\".join(work_path)

        pythoncom.CoInitialize()
        excel = win32.Dispatch('Excel.Application')
        workbook = excel.Workbooks.Open(work_path)
        workbook.PrintOut()
        workbook.Close(SaveChanges=False)
        excel.Quit()
        pythoncom.CoUninitialize()


if __name__ == "__main__":
    ...
